from __future__ import annotations

import csv
import json
from pathlib import Path
from typing import Any, Iterable, Mapping, Sequence

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

FORMULA_PREFIXES = ("=", "+", "-", "@", "\t", "\r")


def collect_fieldnames(records: Iterable[Mapping[str, Any]]) -> list[str]:
    names: list[str] = []
    seen: set[str] = set()
    for record in records:
        for key in record:
            if key not in seen:
                seen.add(key)
                names.append(key)
    return names


def flatten_value(value: Any) -> Any:
    if value is None:
        return ""
    if isinstance(value, (list, tuple, set)):
        return " | ".join(str(v) for v in value)
    if isinstance(value, dict):
        return json.dumps(value, ensure_ascii=False, sort_keys=True)
    return value


def write_csv(
    records: Iterable[Mapping[str, Any]],
    path: str | Path,
    fieldnames: Sequence[str] | None = None,
    escape_formulas: bool = False,
    encoding: str = "utf-8-sig",
) -> Path:
    rows = list(records)
    columns = list(fieldnames or collect_fieldnames(rows))
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding=encoding) as handle:
        writer = csv.DictWriter(handle, fieldnames=columns, extrasaction="ignore")
        writer.writeheader()
        for row in rows:
            out = {}
            for key in columns:
                value = flatten_value(row.get(key))
                if escape_formulas and isinstance(value, str) and value.startswith(FORMULA_PREFIXES):
                    value = "'" + value
                out[key] = value
            writer.writerow(out)
    return path


def xlsx_safe(value: Any) -> Any:
    value = flatten_value(value)
    if isinstance(value, str):
        value = ILLEGAL_CHARACTERS_RE.sub("", value)
    return value


def write_xlsx(
    records: Iterable[Mapping[str, Any]],
    path: str | Path,
    fieldnames: Sequence[str] | None = None,
    sheet_name: str = "data",
) -> Path:
    rows = list(records)
    columns = list(fieldnames or collect_fieldnames(rows))
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = sheet_name[:31]
    sheet.append(columns)
    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
    widths = [len(str(c)) for c in columns]
    for row in rows:
        values = [xlsx_safe(row.get(key)) for key in columns]
        sheet.append(values)
        for index, cell in enumerate(sheet[sheet.max_row]):
            if isinstance(cell.value, str):
                if cell.value.startswith("="):
                    cell.data_type = "s"
                widths[index] = max(widths[index], len(cell.value))
            elif cell.value is not None:
                widths[index] = max(widths[index], len(str(cell.value)))
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = min(max(width + 2, 8), 60)
    sheet.freeze_panes = "A2"
    if columns:
        sheet.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{max(sheet.max_row, 1)}"
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
    return path


def export_records(records: Iterable[Mapping[str, Any]], path: str | Path, **kwargs) -> Path:
    suffix = Path(path).suffix.lower()
    if suffix == ".xlsx":
        return write_xlsx(records, path, **kwargs)
    if suffix in (".csv", ".txt"):
        return write_csv(records, path, **kwargs)
    raise ValueError(f"unsupported export format {suffix!r}; use .csv or .xlsx")
